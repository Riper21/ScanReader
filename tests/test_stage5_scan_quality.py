# -*- coding: utf-8 -*-
"""
Tests for scan-quality remediation (S-1..S-8), по мотивам разбора кейса 4-4-745x1024.jpg:
- S-1: апскейл низких сканов (< 1500 px) в 2x LANCZOS перед VLM.
- S-2: связность статусов: ZT error -> validation.passed=False, quality != excellent;
       OCR_LOW_CONFIDENCE -> штраф 10 п.п. к качеству.
- S-3: нормализация ИП-номера (восстановление слэшей NNNNN/NN/NNNNN-ИП).
- S-4: Excel «Средний балл качества» = quality_score_percent, а не Guardrails.
- S-5: казначейские БИКи 01xxxxxxx валидны.
- S-6: счета Банка России (ГРКЦ) — warning вместо ложного INVALID_BANK_ACCOUNT.
- S-7/S-8: правила ip_number_format / ip_number_format_optional в Guardrails.
"""

from unittest.mock import patch

import pytest
from PIL import Image

from scan_reader.core.utils import normalize_ip_number
from scan_reader.verifier.checksums import validate_bik
from scan_reader.verifier.auditor import ZeroTrustAuditor
from scan_reader.verifier.status import VerificationStatus


# ---------------------------------------------------------------------------
# S-3: нормализация ИП-номера
# ---------------------------------------------------------------------------
def test_s3_normalize_ip_number_restores_slashes():
    """S-3: «107011630001-ИП» -> «10701/16/30001-ИП» (цифры верны, слэши потеряны VLM)."""
    assert normalize_ip_number("107011630001") == "107011630001"  # без -ИП не трогаем
    assert normalize_ip_number("107011630001-ИП") == "10701/16/30001-ИП"
    assert normalize_ip_number("№ 107011630001-ИП") == "10701/16/30001-ИП"
    # 11 цифр (5/2/4): формат кейса «10701/16/3001-ИП»
    assert normalize_ip_number("10701163001-ИП") == "10701/16/3001-ИП"
    # Уже корректный формат не меняется
    assert normalize_ip_number("10701/16/3001-ИП") == "10701/16/3001-ИП"
    assert normalize_ip_number("№ 64947/18/23023-ИП") == "№ 64947/18/23023-ИП"
    # Не-ИП строки не трогаются
    assert normalize_ip_number("ФС006655763") == "ФС006655763"
    assert normalize_ip_number("Вх. № 514") == "Вх. № 514"
    # 14+ цифр (не формат ИП) не трогаются
    assert normalize_ip_number("32230001160010701004") == "32230001160010701004"


def test_s3_salary_schema_restores_ip_number():
    """S-3: Pydantic-схема зарплатных постановлений восстанавливает формат."""
    from scan_reader.doc_types.salary_deductions.schema import SalaryDeductionDoc

    doc = SalaryDeductionDoc.model_validate({"ip_number": "107011630001-ИП", "doc_number": "107011630001-ИП"})
    assert doc.ip_number == "10701/16/30001-ИП"
    assert doc.doc_number == "10701/16/30001-ИП"


def test_s3_enforcement_schema_restores_ip_number():
    """S-3: Схема приказов о возбуждении ИП восстанавливает формат."""
    from scan_reader.doc_types.enforcement_orders.schema import EnforcementOrderDoc

    doc = EnforcementOrderDoc.model_validate({"ip_number": "649471823023-ИП"})
    assert doc.ip_number == "64947/18/23023-ИП"


# ---------------------------------------------------------------------------
# S-5: казначейские БИКи
# ---------------------------------------------------------------------------
def test_s5_treasury_bik_valid():
    """S-5: БИК УФК 015004950 валиден (раньше — ложный INVALID_BIK)."""
    ok, msg = validate_bik("015004950")
    assert ok is True


# ---------------------------------------------------------------------------
# S-6: счета Банка России (ГРКЦ) — warning, не error
# ---------------------------------------------------------------------------
def test_s6_bank_of_russia_account_is_warning_not_error():
    """S-6: счет ГРКЦ Банка России не проходит стандартный ключ -> warning, не error."""
    payload = {
        "payment_details": {
            "payment_account": "40302810200001000046",
            "bik": "041203001",
            "recipient": "ГРКЦ ГУ Банка России по Астраханской обл. г. Астрахань",
        }
    }
    report = ZeroTrustAuditor.audit_document(payload, doc_type="salary_deductions")
    codes = [(i.severity, i.code) for i in report.issues]
    assert ("error", "INVALID_BANK_ACCOUNT") not in codes
    assert ("warning", "BANK_ACCOUNT_UNVERIFIED") in codes
    # Ошибок нет -> не discrepancy (warning не валит документ)
    assert report.status != VerificationStatus.DISCREPANCY_DETECTED


def test_s6_regular_invalid_account_still_error():
    """S-6: обычный невалидный счет (не БР) по-прежнему ошибка."""
    payload = {
        "payment_details": {
            "payment_account": "40702810000000000000",
            "bik": "044525225",
            "recipient": "ООО Ромашка",
        }
    }
    report = ZeroTrustAuditor.audit_document(payload, doc_type="salary_deductions")
    codes = [(i.severity, i.code) for i in report.issues]
    assert ("error", "INVALID_BANK_ACCOUNT") in codes
    assert report.status == VerificationStatus.DISCREPANCY_DETECTED


# ---------------------------------------------------------------------------
# S-2: связность статусов
# ---------------------------------------------------------------------------
def _make_zt(error=False, low_dpi=False):
    """Фабрика VerificationReport-заменителей."""
    from scan_reader.verifier.status import VerificationIssue, VerificationReport

    issues = []
    if error:
        issues.append(VerificationIssue("error", "INVALID_INN", "bad"))
    if low_dpi:
        issues.append(VerificationIssue("warning", "OCR_LOW_CONFIDENCE", "low dpi"))
    details = {"scan_low_quality": low_dpi}
    status = VerificationStatus.DISCREPANCY_DETECTED if error else VerificationStatus.ZERO_TRUST_VERIFIED
    return VerificationReport(status=status, is_valid=not error, issues=issues, details=details)


def test_s2_zt_error_fails_validation_and_quality():
    """S-2: ошибка Zero-Trust -> validation.passed=False, quality_status=needs_attention."""
    from scan_reader.facade import LegalDocPlatformFacade

    validation = {"passed": True, "score": 100.0, "issues": []}
    v, score, status = LegalDocPlatformFacade._combine_quality_and_validation(
        validation, 95.54, "excellent", _make_zt(error=True)
    )
    assert v["passed"] is False
    assert any(i["field"] == "zero_trust" for i in v["issues"])
    assert status == "needs_attention"


def test_s2_low_dpi_penalizes_quality():
    """S-2: низкий DPI -> штраф 10 п.п. и запрет «excellent»."""
    from scan_reader.facade import LegalDocPlatformFacade

    validation = {"passed": True, "score": 100.0, "issues": []}
    v, score, status = LegalDocPlatformFacade._combine_quality_and_validation(
        validation, 95.54, "excellent", _make_zt(low_dpi=True)
    )
    assert score == 85.54
    assert status == "high"
    assert v["passed"] is True
    assert any(i["field"] == "scan" for i in v["issues"])


def test_s2_clean_report_unchanged():
    """S-2: чистый ZT-отчет не меняет метрики."""
    from scan_reader.facade import LegalDocPlatformFacade

    validation = {"passed": True, "score": 100.0, "issues": []}
    v, score, status = LegalDocPlatformFacade._combine_quality_and_validation(
        validation, 95.54, "excellent", _make_zt()
    )
    assert v["passed"] is True
    assert score == 95.54
    assert status == "excellent"


# ---------------------------------------------------------------------------
# S-7/S-8: правила ip_number_format в Guardrails
# ---------------------------------------------------------------------------
def test_s7_ip_number_format_rule():
    """S-7: правило ip_number_format ловит слитные/короткие номера ИП."""
    from scan_reader.facade import LegalDocPlatformFacade

    facade = LegalDocPlatformFacade(enable_cache=False)
    with patch.dict(facade.registry.get("salary_deductions").autonomous_config, {"fields": [
        {"field": "ip_number", "rule": "ip_number_format", "severity": "HIGH", "message": "формат"}
    ]}):
        bad = facade.validate_document({"ip_number": "107011630001"}, "salary_deductions")
        empty = facade.validate_document({"ip_number": ""}, "salary_deductions")
        ok = facade.validate_document({"ip_number": "10701/16/3001-ИП"}, "salary_deductions")

    assert bad["passed"] is False
    assert empty["passed"] is False
    assert ok["passed"] is True


def test_s8_ip_number_format_optional_rule():
    """S-8: опциональный вариант пропускает пустой номер (заявления о возбуждении ИП)."""
    from scan_reader.facade import LegalDocPlatformFacade

    facade = LegalDocPlatformFacade(enable_cache=False)
    with patch.dict(facade.registry.get("enforcement_orders").autonomous_config, {"fields": [
        {"field": "ip_number", "rule": "ip_number_format_optional", "severity": "MEDIUM", "message": "формат"}
    ]}):
        empty_ok = facade.validate_document({"ip_number": ""}, "enforcement_orders")
        bad = facade.validate_document({"ip_number": "107011630001"}, "enforcement_orders")
        ok = facade.validate_document({"ip_number": "64947/18/23023-ИП"}, "enforcement_orders")

    assert empty_ok["passed"] is True
    assert bad["passed"] is False
    assert ok["passed"] is True


def test_s7_salary_guardrails_config_real():
    """S-7: реальный autonomous.json зарплатных постановлений отклоняет слитный номер ИП."""
    from scan_reader.facade import LegalDocPlatformFacade

    facade = LegalDocPlatformFacade(enable_cache=False)
    res = facade.validate_document(
        {"debtor_name": "Иванов И.И.", "claimant_name": "Петров П.П.",
         "finances": {"total_deduction_rub": 1000.0},
         "ip_number": "107011630001", "base_doc_number": ""},
        "salary_deductions",
    )
    assert res["passed"] is False
    messages = " ".join(i.get("message", "") for i in res["issues"])
    assert "NNNNN/NN/NNNNN" in messages


# ---------------------------------------------------------------------------
# S-1: апскейл низких сканов
# ---------------------------------------------------------------------------
def test_s1_small_scan_upscaled():
    """S-1: скан 745x1024 (96 DPI) апскейлится 2x до читаемого размера."""
    from scan_reader.file_processor import FileProcessor

    proc = FileProcessor(max_dimension=2048, min_dimension=1500)
    img = Image.new("RGB", (745, 1024), color="white")
    resized = proc.resize_image_if_needed(img)
    w, h = resized.size
    assert (w, h) == (1490, 2048)


def test_s1_medium_scan_untouched():
    """S-1: изображение в диапазоне [min_dimension, max_dimension] не масштабируется."""
    from scan_reader.file_processor import FileProcessor

    proc = FileProcessor(max_dimension=2048, min_dimension=1500)
    img = Image.new("RGB", (1600, 2000), color="white")
    resized = proc.resize_image_if_needed(img)
    assert resized.size == (1600, 2000)


def test_s1_large_scan_downscaled():
    """S-1: большое изображение по-прежнему уменьшается до лимита."""
    from scan_reader.file_processor import FileProcessor

    proc = FileProcessor(max_dimension=2048, min_dimension=1500)
    img = Image.new("RGB", (4000, 3000), color="white")
    resized = proc.resize_image_if_needed(img)
    assert resized.size == (2048, 1536)


def test_s1_pipeline_prepares_upscaled_uri():
    """S-1: конвейер (process_image_file) отдаёт апскейленный скан в VLM."""
    import tempfile
    import os

    from scan_reader.file_processor import FileProcessor

    proc = FileProcessor(max_dimension=2048, min_dimension=1500)
    img = Image.new("RGB", (745, 1024), color="white")
    with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tf:
        img.save(tf.name, format="JPEG", dpi=(96, 96))
        tmp = tf.name
    try:
        uris = proc.process_image_file(tmp)
        assert len(uris) == 1
        assert uris[0].startswith("data:image/jpeg;base64,")
    finally:
        os.remove(tmp)


# ---------------------------------------------------------------------------
# S-4: Excel «Средний балл качества» = quality_score_percent
# ---------------------------------------------------------------------------
def test_s4_excel_quality_column_uses_quality_score(tmp_path):
    """S-4: сводка Excel показывает quality_score_percent (85.5), а не Guardrails (100)."""
    from scan_reader.excel_exporter import LegalExcelExporter

    items = [{
        "file_name": "doc.pdf",
        "doc_type": "salary_deductions",
        "validation": {"passed": True, "score": 100.0, "issues": []},
        "quality_score_percent": 85.5,
        "quality_status": "high",
        "zero_trust_status": "zero_trust_verified",
        "data": {"finances": {"total_deduction_rub": 100.0}},
    }]
    exporter = LegalExcelExporter(output_dir=str(tmp_path))
    out = exporter.export_results_to_excel(items, output_dir=str(tmp_path))

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb["Сводка"]
    # Строка категории «Взыскание на зарплату»: колонка 6 = «Средний балл качества»
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] and ("заработную" in str(row[0]) or "взыскани" in str(row[0]).lower()):
            assert "85.5%" in str(row[5])
            break
    else:
        pytest.fail("Строка категории salary_deductions не найдена в сводке")


def test_s4_excel_detail_sheet_has_quality_column(tmp_path):
    """S-4: детальный лист содержит колонку «Качество (%)» с реальным значением."""
    from scan_reader.excel_exporter import LegalExcelExporter

    items = [{
        "file_name": "doc.pdf",
        "doc_type": "salary_deductions",
        "validation": {"passed": True, "score": 100.0, "issues": []},
        "quality_score_percent": 85.5,
        "quality_status": "high",
        "zero_trust_status": "zero_trust_verified",
        "data": {"finances": {"total_deduction_rub": 100.0}},
    }]
    exporter = LegalExcelExporter(output_dir=str(tmp_path))
    out = exporter.export_results_to_excel(items, output_dir=str(tmp_path))

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb["Взыскание на зарплату"]
    headers = [c.value for c in ws[1]]
    assert "Качество (%)" in headers
    quality_col = headers.index("Качество (%)") + 1
    assert ws.cell(row=2, column=quality_col).value == 85.5
