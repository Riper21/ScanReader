# -*- coding: utf-8 -*-
"""
Тесты генератора Excel отчетов (excel_exporter.py).
"""

import os
import tempfile
import openpyxl
from scan_reader.excel_exporter import LegalExcelExporter


def test_excel_export_creates_valid_workbook(sample_enforcement_order_data):
    with tempfile.TemporaryDirectory() as tmp_dir:
        exporter = LegalExcelExporter(output_dir=tmp_dir)

        test_records = [
            {
                "file_name": "test1.jpg",
                "doc_type": "enforcement_orders",
                "data": sample_enforcement_order_data,
                "validation": {"passed": True, "score": 100.0, "issues": []}
            },
            {
                "file_name": "test2.jpg",
                "doc_type": "enforcement_orders",
                "data": sample_enforcement_order_data,
                "validation": {"passed": False, "score": 66.7, "issues": [{"field": "claimant.name"}]}
            }
        ]

        out_path = exporter.export_results_to_excel(test_records, file_name="test_report.xlsx")
        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        sheet_names = wb.sheetnames
        assert "Сводка" in sheet_names
        assert "Приказы ИП" in sheet_names

        ws_sum = wb["Сводка"]
        assert ws_sum["A1"].value == "📊 СВОДНЫЙ ДАШБОРД РАСПОЗНАВАНИЯ ДОКУМЕНТОВ"
        wb.close()


def test_excel_export_includes_zero_trust_status(sample_enforcement_order_data):
    with tempfile.TemporaryDirectory() as tmp_dir:
        exporter = LegalExcelExporter(output_dir=tmp_dir)

        test_records = [
            {
                "file_name": "verified.jpg",
                "doc_type": "enforcement_orders",
                "data": sample_enforcement_order_data,
                "validation": {"passed": True, "score": 100.0, "issues": []},
                "zero_trust_status": "zero_trust_verified"
            },
            {
                "file_name": "discrepancy.jpg",
                "doc_type": "enforcement_orders",
                "data": sample_enforcement_order_data,
                "validation": {"passed": False, "score": 50.0, "issues": []},
                "zero_trust_status": "discrepancy_detected"
            }
        ]

        out_path = exporter.export_results_to_excel(test_records, file_name="zt_report.xlsx")
        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        ws_detail = wb["Приказы ИП"]

        # Check that headers include 'Zero-Trust Статус'
        header_vals = [cell.value for cell in ws_detail[1]]
        assert "Zero-Trust Статус" in header_vals

        # Check status values rendered in rows
        zt_col_idx = header_vals.index("Zero-Trust Статус") + 1
        row2_val = ws_detail.cell(row=2, column=zt_col_idx).value
        row3_val = ws_detail.cell(row=3, column=zt_col_idx).value

        assert "Zero-Trust" in row2_val
        assert "Расхождение" in row3_val
        wb.close()

