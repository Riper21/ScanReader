# -*- coding: utf-8 -*-
"""
Модуль профессионального экспорта реестров распознанных документов в Excel (.xlsx).
Создает стилизованные многостраничные книги со сводным дашбордом,
цветовой индикацией типов документов и статусов валидации.
"""

import os
import tempfile
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    openpyxl = None
    Font = PatternFill = Alignment = Border = Side = None
    get_column_letter = None
    OPENPYXL_AVAILABLE = False

from .core.utils import get_logger
from .type_registry import get_registry

logger = get_logger("excel_exporter")


class LegalExcelExporter:
    """
    Генератор Excel-отчетов и сводных реестров на базе openpyxl.
    """

    def __init__(self, output_dir: Optional[str] = None):
        self.output_dir: str = ""
        if output_dir:
            self.output_dir = output_dir
        else:
            env_out = os.getenv("SCANREADER_OUTPUT_DIR")
            if env_out:
                self.output_dir = env_out
            else:
                legacy_results = os.path.join(os.getcwd(), "Результаты")
                default_output = os.path.join(os.getcwd(), "output")
                if os.path.exists(legacy_results) and not os.path.exists(default_output):
                    self.output_dir = legacy_results
                else:
                    self.output_dir = default_output
        os.makedirs(self.output_dir, exist_ok=True)
        self.registry = get_registry()

    def export_results_to_excel(
        self,
        results: List[Dict[str, Any]],
        file_name: str = "Сводный_реестр_документов.xlsx",
        output_dir: Optional[str] = None
    ) -> str:
        """
        Экспортирует список обработанных документов в многостраничный Excel файл
        с полной цветовой индикацией статусов Zero-Trust и Guardrails.
        """
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("Для экспорта в Excel установите openpyxl: `py -3 -m pip install openpyxl`")

        target_dir = output_dir or self.output_dir
        os.makedirs(target_dir, exist_ok=True)

        wb = openpyxl.Workbook()
        # Удаляем дефолтный лист
        wb.remove(wb.active)

        # Стили
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        font_bold = Font(name="Calibri", size=10, bold=True)
        
        fill_dark_blue = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        fill_header_order = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
        fill_header_exec = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        fill_header_salary = PatternFill(start_color="385723", end_color="385723", fill_type="solid")

        fill_valid_pass = PatternFill(start_color="F2F9F1", end_color="F2F9F1", fill_type="solid")
        fill_valid_fail = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")

        # Zero-Trust Color Badges
        fill_zt_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        font_zt_green = Font(name="Calibri", size=10, bold=True, color="276A3C")

        fill_zt_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        font_zt_red = Font(name="Calibri", size=10, bold=True, color="C00000")

        fill_zt_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        font_zt_yellow = Font(name="Calibri", size=10, bold=True, color="7F6000")

        fill_zt_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        font_zt_gray = Font(name="Calibri", size=10, color="595959")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")

        # 1. Лист "Сводка" (Summary)
        ws_sum = wb.create_sheet(title="Сводка")
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum.merge_cells("A1:F1")
        ws_sum["A1"] = "📊 СВОДНЫЙ ДАШБОРД РАСПОЗНАВАНИЯ ДОКУМЕНТОВ"
        ws_sum["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        ws_sum["A1"].fill = fill_dark_blue
        ws_sum["A1"].alignment = align_center
        ws_sum.row_dimensions[1].height = 35

        headers_sum = [
            "Тип документа",
            "Количество",
            "Валидация (Guardrails + Zero-Trust)",
            "Zero-Trust Проверено",
            "Сумма по реестру (руб)",
            "Средний балл качества"
        ]

        for col_idx, h_text in enumerate(headers_sum, 1):
            cell = ws_sum.cell(row=3, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_dark_blue
            cell.alignment = align_center
            cell.border = thin_border
        ws_sum.row_dimensions[3].height = 25

        # Группировка результатов по doc_type
        by_type: Dict[str, List[Dict[str, Any]]] = {}
        for r in results:
            dt = r.get("doc_type", "unknown")
            if dt not in by_type:
                by_type[dt] = []
            by_type[dt].append(r)

        sum_row = 4
        total_all_count = len(results)
        total_all_passed = 0
        total_all_zt = 0
        total_all_sum = 0.0

        for plugin in self.registry.enabled().values():
            pid = plugin.id
            items = by_type.get(pid, [])
            count = len(items)
            passed = sum(1 for i in items if i.get("validation", {}).get("passed", False))
            zt_verified = sum(
                1 for i in items
                if (i.get("zero_trust_status") == "zero_trust_verified"
                    or i.get("zero_trust", {}).get("status") == "zero_trust_verified")
            )

            total_all_passed += passed
            total_all_zt += zt_verified

            # Сумма
            total_sum = 0.0
            for i in items:
                fin = i.get("data", {}).get("finances", {})
                if isinstance(fin, dict):
                    val = fin.get("total_rub") or fin.get("total_deduction_rub") or fin.get("debt_amount_rub")
                    if val and isinstance(val, (int, float)):
                        total_sum += float(val)
            total_all_sum += total_sum

            # Средний балл качества — из quality_score_percent (S-4: ранее брался
            # validation.score от Guardrails, что показывало 100% при реальном 95.4%)
            avg_score = (
                round(sum(float(i.get("quality_score_percent", 100.0)) for i in items) / count, 1)
                if count > 0 else 100.0
            )

            ws_sum.cell(row=sum_row, column=1, value=plugin.title).font = font_bold
            ws_sum.cell(row=sum_row, column=2, value=count).alignment = align_center
            ws_sum.cell(row=sum_row, column=3, value=f"{passed}/{count}").alignment = align_center
            ws_sum.cell(row=sum_row, column=4, value=f"{zt_verified}/{count}").alignment = align_center
            ws_sum.cell(row=sum_row, column=5, value=total_sum).number_format = "#,##0.00"
            ws_sum.cell(row=sum_row, column=6, value=f"{avg_score}%").alignment = align_center

            for c in range(1, 7):
                ws_sum.cell(row=sum_row, column=c).border = thin_border
            ws_sum.row_dimensions[sum_row].height = 20
            sum_row += 1

        # Итоговая строка
        ws_sum.cell(row=sum_row, column=1, value="ИТОГО:").font = font_bold
        ws_sum.cell(row=sum_row, column=2, value=total_all_count).font = font_bold
        ws_sum.cell(row=sum_row, column=2).alignment = align_center
        ws_sum.cell(row=sum_row, column=3, value=f"{total_all_passed}/{total_all_count}").alignment = align_center
        ws_sum.cell(row=sum_row, column=4, value=f"{total_all_zt}/{total_all_count}").alignment = align_center
        ws_sum.cell(row=sum_row, column=5, value=total_all_sum).number_format = "#,##0.00"
        ws_sum.cell(row=sum_row, column=6, value="").alignment = align_center
        for c in range(1, 7):
            ws_sum.cell(row=sum_row, column=c).border = thin_border
            ws_sum.cell(row=sum_row, column=c).font = font_bold
        ws_sum.row_dimensions[sum_row].height = 22

        # 2. Детальные листы по каждому типу документа
        for plugin in self.registry.enabled().values():
            pid = plugin.id
            items = by_type.get(pid, [])
            if not items:
                continue

            ws = wb.create_sheet(title=plugin.short_title[:31])
            ws.views.sheetView[0].showGridLines = True

            # Выбор цвета шапки
            fill_header = fill_dark_blue
            if "enforcement" in pid:
                fill_header = fill_header_order
            elif "executive" in pid:
                fill_header = fill_header_exec
            elif "salary" in pid:
                fill_header = fill_header_salary

            cols = [{"path": "file_name", "label": "Имя файла", "kind": "text"}] + plugin.flat_columns + [
                {"path": "validation.score", "label": "Guardrails (%)", "kind": "number"},
                {"path": "quality_score_percent", "label": "Качество (%)", "kind": "number"},
                {"path": "zero_trust_status", "label": "Zero-Trust Статус", "kind": "zt_status"}
            ]

            # Заголовки
            for col_idx, col_def in enumerate(cols, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_def.get("label", ""))
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
                cell.border = thin_border
            ws.row_dimensions[1].height = 28

            def _get_field(data_dict: Dict[str, Any], path_str: str) -> Any:
                parts = path_str.split(".")
                curr: Any = data_dict
                for p in parts:
                    if isinstance(curr, dict):
                        curr = curr.get(p)
                    else:
                        return ""
                return curr if curr is not None else ""

            # Данные
            for row_idx, item in enumerate(items, 2):
                data_obj = item.get("data", {})
                is_valid = item.get("validation", {}).get("passed", True)
                zt_raw = (
                    item.get("zero_trust_status")
                    or item.get("zero_trust", {}).get("status")
                    or "vlm_unverified"
                )

                for col_idx, col_def in enumerate(cols, 1):
                    p_path = col_def.get("path", "")
                    if p_path.startswith("validation."):
                        val = item.get("validation", {}).get(p_path.replace("validation.", ""), "")
                    elif p_path == "file_name":
                        val = item.get("file_name", "")
                    elif p_path == "quality_score_percent":
                        val = item.get("quality_score_percent", "")
                    elif p_path == "zero_trust_status":
                        val = zt_raw
                    else:
                        val = _get_field(data_obj, p_path)

                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.font = font_data

                    if col_def.get("kind") == "number":
                        try:
                            if val != "":
                                cell.value = float(val)
                                cell.number_format = "#,##0.00"
                            else:
                                cell.value = ""
                        except Exception:
                            cell.value = str(val)
                        cell.alignment = align_right
                    elif col_def.get("kind") == "boolean":
                        cell.value = "Да" if val else "Нет"
                        cell.alignment = align_center
                    elif col_def.get("kind") == "zt_status":
                        # Zero-Trust Badge
                        zt_str = str(val).lower()
                        if zt_str == "zero_trust_verified":
                            cell.value = "✔ Zero-Trust"
                            cell.fill = fill_zt_green
                            cell.font = font_zt_green
                        elif zt_str == "discrepancy_detected":
                            cell.value = "✘ Расхождение"
                            cell.fill = fill_zt_red
                            cell.font = font_zt_red
                        elif zt_str == "heuristic_fallback":
                            cell.value = "⚠ Эвристика"
                            cell.fill = fill_zt_yellow
                            cell.font = font_zt_yellow
                        elif zt_str == "rejected_unsupported":
                            cell.value = "⛔ Не поддерж."
                            cell.fill = fill_zt_red
                            cell.font = font_zt_red
                        else:
                            cell.value = "🔍 VLM анализ"
                            cell.fill = fill_zt_gray
                            cell.font = font_zt_gray
                        cell.alignment = align_center
                    else:
                        cell.value = str(val) if val is not None else ""
                        cell.alignment = align_left

                # Подсветка строки (мягкий оттенок, если колонка не имеет собственного бейджа)
                row_fill = fill_valid_pass if (is_valid and zt_raw != "discrepancy_detected") else fill_valid_fail
                for col_idx in range(1, len(cols)):
                    ws.cell(row=row_idx, column=col_idx).fill = row_fill
                ws.row_dimensions[row_idx].height = 20

            # Автоподбор ширины колонок
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Автоподбор ширины на листе Сводка
        for col in ws_sum.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 15)

        out_path = os.path.join(target_dir, file_name)
        # M-19: атомарное сохранение Excel (временный файл в целевом каталоге + os.replace)
        tmp_fd, tmp_path = tempfile.mkstemp(prefix=f".{file_name}.tmp-", suffix=".xlsx", dir=target_dir)
        os.close(tmp_fd)
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, out_path)
        except Exception:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError as e:
                    logger.debug(f"Не удалось удалить временный Excel файл '{tmp_path}': {e}")
            raise
        logger.info(f"📊 Реестр успешно экспортирован в Excel: {out_path}")
        return out_path
