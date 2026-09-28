# -*- coding: utf-8 -*-
"""
Модуль комплексной метрической оценки качества распознавания юридических документов для ScanReader.
Рассчитывает процентную точность (% Quality Score / Accuracy / Completeness):
1. В режиме бенчмарка (Benchmark vs Ground Truth): сравнение AI-извлечения с эталоном.
2. В автономном режиме (Autonomous Quality & Completeness & Guardrails Validation):
   оценка полноты заполнения схемы, валидности форматов (даты, номера дел, бланки),
   проверка контрольных разрядов ИНН (ФНС) и балансовой непротиворечивости сумм.

Формирует:
- JSON (.json) структурированные метрические файлы для API и интеграций:
  - run_metrics_summary.json (сводка запуска)
  - {doc_type}_quality_metrics.json (детализация по категории)
  - metrics_history.json (накопительная история качества)
- Markdown (.md) сводный отчет run_metrics_summary.md
- Excel (.xlsx) сводный отчет run_metrics_summary.xlsx (при наличии openpyxl)
"""

import os
import re
import json
import difflib
import datetime
from typing import Dict, Any, List, Optional, Tuple

from .utils import get_logger
from .io_utils import write_atomic

logger = get_logger("core.metrics_evaluator")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ==============================================================================
# 1. МАТЕМАТИЧЕСКИЕ И ТЕКСТОВЫЕ МЕТРИКИ СХОДСТВА (BENCHMARK MODE)
# ==============================================================================

def normalize_text(text: Optional[Any]) -> str:
    """Удаляет лишние пробелы, переносы строк и приводит к нижнему регистру для нечувствительного сравнения."""
    if text is None:
        return ""
    text = str(text).lower().strip()
    text = re.sub(r"\s+", " ", text)
    replacements = {
        "№": "no", "«": '"', "»": '"', "—": "-", "–": "-",
        "а": "a", "в": "b", "е": "e", "к": "k", "м": "m", "н": "h", "о": "o", "р": "p", "с": "c", "т": "t", "х": "x"
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return text


def exact_match_score(pred: Any, gt: Any) -> float:
    """Оценка 100% при точном совпадении после нормализации, иначе 0%."""
    n_pred = normalize_text(pred)
    n_gt = normalize_text(gt)
    if not n_gt and not n_pred:
        return 100.0
    if not n_gt or not n_pred:
        return 0.0
    return 100.0 if n_pred == n_gt else 0.0


def text_similarity_score(pred: Optional[Any], gt: Optional[Any]) -> float:
    """Расчет схожести строк (SequenceMatcher Ratio) в диапазоне 0.0 - 100.0%."""
    n_pred = normalize_text(pred)
    n_gt = normalize_text(gt)
    if not n_gt and not n_pred:
        return 100.0
    if not n_gt or not n_pred:
        return 0.0
    if n_pred == n_gt:
        return 100.0
    ratio = difflib.SequenceMatcher(None, n_pred, n_gt).ratio()
    return round(ratio * 100.0, 2)


def numeric_proximity_score(pred: Optional[float], gt: Optional[float], tolerance: float = 0.01) -> float:
    """Оценка числовой близости сумм в диапазоне 0.0 - 100.0%."""
    if gt is None and pred is None:
        return 100.0
    if gt is None or pred is None:
        return 0.0
    try:
        p_val = float(pred)
        g_val = float(gt)
    except (ValueError, TypeError):
        return 0.0

    if abs(p_val - g_val) <= tolerance:
        return 100.0
    diff = abs(p_val - g_val)
    denom = max(abs(g_val), 1.0)
    score = max(0.0, 1.0 - (diff / denom)) * 100.0
    return round(score, 2)


# ==============================================================================
# 2. АВТОНОМНЫЕ ВАЛИДАТОРЫ И ПРОВЕРКИ РЕКВИЗИТОВ (GUARDRAILS)
# ==============================================================================

def validate_inn_string(details_str: Optional[str]) -> Tuple[bool, Optional[str], float]:
    """Проверяет наличие корректного ИНН (10 или 12 знаков) по контрольным разрядам ФНС."""
    if not details_str:
        return False, None, 0.0
    matches = re.findall(r"\b(\d{10}|\d{12})\b", str(details_str))
    if not matches:
        return False, None, 50.0  # Реквизиты есть, но без ИНН

    for inn in matches:
        digits = [int(c) for c in inn]
        if len(digits) == 10:
            coeffs = [2, 4, 10, 3, 5, 9, 4, 6, 8]
            chk = sum(d * c for d, c in zip(digits[:9], coeffs)) % 11 % 10
            if chk == digits[9]:
                return True, inn, 100.0
        elif len(digits) == 12:
            coeffs1 = [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
            chk1 = sum(d * c for d, c in zip(digits[:10], coeffs1)) % 11 % 10
            coeffs2 = [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
            chk2 = sum(d * c for d, c in zip(digits[:11], coeffs2)) % 11 % 10
            if chk1 == digits[10] and chk2 == digits[11]:
                return True, inn, 100.0
    return True, matches[0], 85.0


def validate_case_number_format(case_num: Optional[str]) -> Tuple[bool, float]:
    """
    Проверяет формат судебного дела (А40-1234/2021, 2-123/2020 и т.д.).
    S-12: длинные номера без разделителей («-», «/») подозрительны —
    VLM мог склеить или исказить номер; оценка снижается до 60.
    """
    if not case_num:
        return False, 0.0
    c_str = str(case_num).strip()
    pattern = r"^[А-Яа-я0-9A-Za-z\s\№\-\–\—\/\.]+$"
    if re.match(pattern, c_str) and len(c_str) >= 4 and any(c.isdigit() for c in c_str):
        if len(c_str) >= 6 and "/" not in c_str and "-" not in c_str and "–" not in c_str:
            return True, 60.0
        return True, 100.0
    return False, 40.0


def validate_ip_number_format(ip_num: Optional[str]) -> Tuple[bool, float]:
    """Проверяет формат исполнительного производства (ХХХХХ/ГГ/ДД/РЕГИОН)."""
    if not ip_num:
        return False, 0.0
    i_str = str(ip_num).strip()
    if "/" in i_str and any(c.isdigit() for c in i_str):
        return True, 100.0
    if len(i_str) >= 4 and any(c.isdigit() for c in i_str):
        return True, 80.0
    return False, 30.0


def validate_date_string(date_str: Optional[str]) -> Tuple[bool, float]:
    """Проверяет дату на соответствие формату ДД.ММ.ГГГГ."""
    if not date_str:
        return False, 0.0
    d_str = str(date_str).strip()
    m = re.search(r"\b(\d{1,2})[\.\/\-](\d{1,2})[\.\/\-](\d{2,4})\b", d_str)
    if m:
        return True, 100.0
    for month in ["январ", "феврал", "март", "апрел", "ма", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр"]:
        if month in d_str.lower():
            return True, 95.0
    return False, 20.0


def validate_math_balance(
    main_debt: Optional[float],
    penalty: Optional[float],
    fee: Optional[float],
    other: Optional[float],
    total: Optional[float]
) -> Tuple[bool, float, float]:
    """
    Проверяет математический баланс:
    Основной долг + Пени + Госпошлина + Прочие = ИТОГО (с допуском 0.05 руб).
    Возвращает (is_balanced, diff_rub, score_percent).
    """
    has_total = total is not None and isinstance(total, (int, float)) and total > 0
    components = [v for v in [main_debt, penalty, fee, other] if v is not None and isinstance(v, (int, float))]

    if not has_total and not components:
        return True, 0.0, 100.0

    if has_total and not components:
        return True, 0.0, 95.0

    calc_sum = sum(c for c in components if c is not None)
    if has_total and total is not None:
        total_f = float(total)
        diff = abs(calc_sum - total_f)
        if diff <= 0.05:
            return True, 0.0, 100.0
        elif diff <= 1.0:
            return True, diff, 90.0
        else:
            denom = max(total_f, 1.0)
            score = max(0.0, 1.0 - (diff / denom)) * 100.0
            return False, round(diff, 2), round(score, 2)
    else:
        return True, 0.0, 90.0


def validate_deduction_rate(rate_str: Optional[str]) -> Tuple[bool, float]:
    """Проверяет процент удержания по 229-ФЗ (25%, 50%, 70%, 1/4)."""
    if not rate_str:
        return False, 0.0
    r_str = str(rate_str).strip()
    if any(k in r_str for k in ["50%", "25%", "70%", "1/4", "1/3", "1/2", "30%", "20%"]):
        return True, 100.0
    if "%" in r_str or "доля" in r_str or "доли" in r_str or "заработк" in r_str:
        return True, 90.0
    return False, 40.0


# ==============================================================================
# 3. ОЦЕНКА ДОКУМЕНТОВ В РЕЖИМЕ BENCHMARK (VS GROUND TRUTH)
# ==============================================================================

def evaluate_executive_document_benchmark(pred: Dict[str, Any], gt: Dict[str, Any]) -> Dict[str, Any]:
    c_pred = pred.get("court") or {}
    c_gt = gt.get("court") or {}
    cl_pred = pred.get("claimant") or {}
    cl_gt = gt.get("claimant") or {}
    db_pred = pred.get("debtor") or {}
    db_gt = gt.get("debtor") or {}
    f_pred = pred.get("finances") or pred.get("financials") or {}
    f_gt = gt.get("finances") or gt.get("financials") or {}

    scores = {
        "Номер дела": exact_match_score(pred.get("case_number"), gt.get("case_number")),
        "Дата судебного акта": exact_match_score(pred.get("act_date"), gt.get("act_date")),
        "Суд (Наименование)": text_similarity_score(c_pred.get("name"), c_gt.get("name")),
        "Суд (Адрес)": text_similarity_score(c_pred.get("address"), c_gt.get("address")),
        "Взыскатель (Наименование/ФИО)": text_similarity_score(cl_pred.get("name"), cl_gt.get("name")),
        "Взыскатель (Реквизиты)": text_similarity_score(cl_pred.get("details"), cl_gt.get("details")),
        "Должник (Наименование/ФИО)": text_similarity_score(db_pred.get("name"), db_gt.get("name")),
        "Должник (Реквизиты)": text_similarity_score(db_pred.get("details"), db_gt.get("details")),
        "Предмет иска": text_similarity_score(pred.get("claim_subject"), gt.get("claim_subject")),
        "Резолюция суда": text_similarity_score(pred.get("decision_summary"), gt.get("decision_summary")),
        "Основной долг": numeric_proximity_score(f_pred.get("main_debt_rub"), f_gt.get("main_debt_rub")),
        "Пени / Проценты": numeric_proximity_score(f_pred.get("interest_penalty_rub"), f_gt.get("interest_penalty_rub")),
        "Госпошлина / Сбор": numeric_proximity_score(f_pred.get("court_fee_rub"), f_gt.get("court_fee_rub")),
        "Итоговая сумма": numeric_proximity_score(f_pred.get("total_rub"), f_gt.get("total_rub")),
        "Серия бланка Гознака": exact_match_score(pred.get("blank_series"), gt.get("blank_series")),
        "Номер бланка Гознака": exact_match_score(pred.get("blank_number"), gt.get("blank_number")),
    }

    weights = {
        "Номер дела": 0.10, "Дата судебного акта": 0.05, "Суд (Наименование)": 0.08, "Суд (Адрес)": 0.02,
        "Взыскатель (Наименование/ФИО)": 0.10, "Взыскатель (Реквизиты)": 0.06,
        "Должник (Наименование/ФИО)": 0.10, "Должник (Реквизиты)": 0.06,
        "Предмет иска": 0.06, "Резолюция суда": 0.06,
        "Основной долг": 0.07, "Пени / Проценты": 0.03, "Госпошлина / Сбор": 0.03, "Итоговая сумма": 0.10,
        "Серия бланка Гознака": 0.04, "Номер бланка Гознака": 0.05,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": gt.get("file_name", pred.get("file_name", "")),
        "mode": "benchmark",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": True
    }


def evaluate_enforcement_document_benchmark(pred: Dict[str, Any], gt: Dict[str, Any]) -> Dict[str, Any]:
    fssp_pred = pred.get("fssp") or pred.get("authority") or {}
    fssp_gt = gt.get("fssp") or gt.get("authority") or {}
    c_pred = pred.get("court") or {}
    c_gt = gt.get("court") or {}
    cl_pred = pred.get("claimant") or {}
    cl_gt = gt.get("claimant") or {}
    db_pred = pred.get("debtor") or {}
    db_gt = gt.get("debtor") or {}
    f_pred = pred.get("finances") or {}
    f_gt = gt.get("finances") or {}

    scores = {
        "Вид документа": text_similarity_score(pred.get("doc_type"), gt.get("doc_type")),
        "Дата документа": exact_match_score(pred.get("doc_date"), gt.get("doc_date")),
        "Номер ИП / Входящий": text_similarity_score(pred.get("ip_number") or pred.get("reg_number"), gt.get("ip_number") or gt.get("reg_number")),
        "Орган ФССП": text_similarity_score(fssp_pred.get("name"), fssp_gt.get("name")),
        "Пристав / Должностное лицо": text_similarity_score(fssp_pred.get("officer"), fssp_gt.get("officer")),
        "Суд-основание": text_similarity_score(c_pred.get("name"), c_gt.get("name")),
        "Номер судебного дела": text_similarity_score(c_pred.get("case_number"), c_gt.get("case_number")),
        "Взыскатель": text_similarity_score(cl_pred.get("name"), cl_gt.get("name")),
        "Реквизиты взыскателя": text_similarity_score(cl_pred.get("details"), cl_gt.get("details")),
        "Должник": text_similarity_score(db_pred.get("name"), db_gt.get("name")),
        "Реквизиты должника": text_similarity_score(db_pred.get("details"), db_gt.get("details")),
        "Предмет исполнения": text_similarity_score(pred.get("claim_subject"), gt.get("claim_subject")),
        "Основной долг": numeric_proximity_score(f_pred.get("main_debt_rub"), f_gt.get("main_debt_rub")),
        "Судебные расходы": numeric_proximity_score(f_pred.get("court_costs_rub"), f_gt.get("court_costs_rub")),
        "Итоговая сумма требований": numeric_proximity_score(f_pred.get("total_rub"), f_gt.get("total_rub")),
    }

    weights = {
        "Вид документа": 0.05, "Дата документа": 0.05, "Номер ИП / Входящий": 0.10, "Орган ФССП": 0.08,
        "Пристав / Должностное лицо": 0.04, "Суд-основание": 0.06, "Номер судебного дела": 0.08,
        "Взыскатель": 0.10, "Реквизиты взыскателя": 0.06, "Должник": 0.10, "Реквизиты должника": 0.06,
        "Предмет исполнения": 0.07, "Основной долг": 0.07, "Судебные расходы": 0.03,
        "Итоговая сумма требований": 0.06,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": gt.get("file_name", pred.get("file_name", "")),
        "mode": "benchmark",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": True
    }


def evaluate_salary_document_benchmark(pred: Dict[str, Any], gt: Dict[str, Any]) -> Dict[str, Any]:
    auth_pred = pred.get("authority") or {}
    auth_gt = gt.get("authority") or {}
    emp_pred = pred.get("employer") or {}
    emp_gt = gt.get("employer") or {}
    f_pred = pred.get("finances") or {}
    f_gt = gt.get("finances") or {}

    scores = {
        "Дата постановления": exact_match_score(pred.get("doc_date"), gt.get("doc_date")),
        "Номер ИП": text_similarity_score(pred.get("ip_number"), gt.get("ip_number")),
        "Орган исполнения": text_similarity_score(auth_pred.get("name"), auth_gt.get("name")),
        "Юрисдикция": text_similarity_score(auth_pred.get("jurisdiction"), auth_gt.get("jurisdiction")),
        "Работодатель": text_similarity_score(emp_pred.get("name"), emp_gt.get("name")),
        "Адрес работодателя": text_similarity_score(emp_pred.get("address"), emp_gt.get("address")),
        "Взыскатель": text_similarity_score(pred.get("claimant_name"), gt.get("claimant_name")),
        "Должник (Сотрудник)": text_similarity_score(pred.get("debtor_name"), gt.get("debtor_name")),
        "Реквизиты должника": text_similarity_score(pred.get("debtor_details"), gt.get("debtor_details")),
        "Предмет взыскания": text_similarity_score(pred.get("claim_subject"), gt.get("claim_subject")),
        "Процент удержания": text_similarity_score(f_pred.get("deduction_percentage"), f_gt.get("deduction_percentage")),
        "Основной долг": numeric_proximity_score(f_pred.get("debt_amount_rub"), f_gt.get("debt_amount_rub")),
        "Исполнительский сбор": numeric_proximity_score(f_pred.get("fee_penalty_rub"), f_gt.get("fee_penalty_rub")),
        "Итого к удержанию": numeric_proximity_score(f_pred.get("total_deduction_rub"), f_gt.get("total_deduction_rub")),
        "Первичный судебный акт": text_similarity_score(pred.get("base_doc"), gt.get("base_doc")),
    }

    weights = {
        "Дата постановления": 0.05, "Номер ИП": 0.10, "Орган исполнения": 0.08, "Юрисдикция": 0.04,
        "Работодатель": 0.08, "Адрес работодателя": 0.04, "Взыскатель": 0.10, "Должник (Сотрудник)": 0.10,
        "Реквизиты должника": 0.05, "Предмет взыскания": 0.06, "Процент удержания": 0.10,
        "Основной долг": 0.06, "Исполнительский сбор": 0.04, "Итого к удержанию": 0.06,
        "Первичный судебный акт": 0.05,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": gt.get("file_name", pred.get("file_name", "")),
        "mode": "benchmark",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": True
    }


# ==============================================================================
# 4. АВТОНОМНАЯ ОЦЕНКА КАЧЕСТВА БЕЗ GROUND TRUTH (AUTONOMOUS QUALITY SCORE)
# ==============================================================================

def evaluate_executive_document_autonomous(doc: Dict[str, Any]) -> Dict[str, Any]:
    c = doc.get("court") or {}
    cl = doc.get("claimant") or {}
    db = doc.get("debtor") or {}
    f = doc.get("finances") or {}

    _, case_score = validate_case_number_format(doc.get("case_number"))
    _, date_score = validate_date_string(doc.get("act_date"))
    _, _, cl_inn_score = validate_inn_string(cl.get("details"))
    _, _, db_inn_score = validate_inn_string(db.get("details"))
    is_bal, _, bal_score = validate_math_balance(
        f.get("main_debt_rub"), f.get("interest_penalty_rub"), f.get("court_fee_rub"), f.get("other_rub"), f.get("total_rub")
    )

    scores = {
        "Номер дела": case_score,
        "Дата судебного акта": date_score,
        "Суд (Наименование)": 100.0 if len(str(c.get("name", "")).strip()) > 3 else (50.0 if c.get("name") else 0.0),
        "Суд (Адрес)": 100.0 if len(str(c.get("address", "")).strip()) > 3 else (60.0 if c.get("address") else 0.0),
        "Взыскатель (Наименование/ФИО)": 100.0 if len(str(cl.get("name", "")).strip()) > 2 else 0.0,
        "Взыскатель (Реквизиты и ИНН)": cl_inn_score,
        "Должник (Наименование/ФИО)": 100.0 if len(str(db.get("name", "")).strip()) > 2 else 0.0,
        "Должник (Реквизиты и ИНН)": db_inn_score,
        "Предмет иска": 100.0 if len(str(doc.get("claim_subject", "")).strip()) > 3 else 0.0,
        "Резолюция суда": 100.0 if len(str(doc.get("decision_summary", "")).strip()) > 3 else 0.0,
        "Основной долг": 100.0 if f.get("main_debt_rub") is not None else 80.0,
        "Пени / Проценты": 100.0 if f.get("interest_penalty_rub") is not None else 90.0,
        "Госпошлина / Сбор": 100.0 if f.get("court_fee_rub") is not None else 90.0,
        "Итоговая сумма и Баланс": bal_score,
        "Серия бланка Гознака": 100.0 if doc.get("blank_series") else 70.0,
        "Номер бланка Гознака": 100.0 if doc.get("blank_number") else 70.0,
    }

    weights = {
        "Номер дела": 0.10, "Дата судебного акта": 0.05, "Суд (Наименование)": 0.08, "Суд (Адрес)": 0.02,
        "Взыскатель (Наименование/ФИО)": 0.10, "Взыскатель (Реквизиты и ИНН)": 0.06,
        "Должник (Наименование/ФИО)": 0.10, "Должник (Реквизиты и ИНН)": 0.06,
        "Предмет иска": 0.06, "Резолюция суда": 0.06,
        "Основной долг": 0.07, "Пени / Проценты": 0.03, "Госпошлина / Сбор": 0.03, "Итоговая сумма и Баланс": 0.10,
        "Серия бланка Гознака": 0.04, "Номер бланка Гознака": 0.05,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": doc.get("file_name", ""),
        "mode": "autonomous",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": bool(is_bal and case_score >= 70.0 and date_score >= 70.0)
    }


def evaluate_enforcement_document_autonomous(doc: Dict[str, Any]) -> Dict[str, Any]:
    fssp = doc.get("fssp") or doc.get("authority") or {}
    c = doc.get("court") or {}
    cl = doc.get("claimant") or {}
    db = doc.get("debtor") or {}
    f = doc.get("finances") or {}

    _, date_score = validate_date_string(doc.get("doc_date"))
    _, ip_score = validate_ip_number_format(doc.get("ip_number") or doc.get("reg_number"))
    _, _, cl_inn_score = validate_inn_string(cl.get("details"))
    _, _, db_inn_score = validate_inn_string(db.get("details"))
    is_bal, _, bal_score = validate_math_balance(
        f.get("main_debt_rub"), f.get("court_costs_rub"), None, None, f.get("total_rub")
    )

    scores = {
        "Вид документа": 100.0 if doc.get("doc_type") else 0.0,
        "Дата документа": date_score,
        "Номер ИП / Входящий": ip_score,
        "Орган ФССП": 100.0 if len(str(fssp.get("name", "")).strip()) > 3 else 0.0,
        "Пристав / Должностное лицо": 100.0 if fssp.get("officer") else 60.0,
        "Суд-основание": 100.0 if c.get("name") else 60.0,
        "Номер судебного дела": 100.0 if c.get("case_number") else 60.0,
        "Взыскатель": 100.0 if len(str(cl.get("name", "")).strip()) > 2 else 0.0,
        "Реквизиты взыскателя": cl_inn_score,
        "Должник": 100.0 if len(str(db.get("name", "")).strip()) > 2 else 0.0,
        "Реквизиты должника": db_inn_score,
        "Предмет исполнения": 100.0 if doc.get("claim_subject") else 50.0,
        "Основной долг": 100.0 if f.get("main_debt_rub") is not None else 80.0,
        "Судебные расходы": 100.0 if f.get("court_costs_rub") is not None else 90.0,
        "Итоговая сумма и Баланс": bal_score,
    }

    weights = {
        "Вид документа": 0.05, "Дата документа": 0.06, "Номер ИП / Входящий": 0.10, "Орган ФССП": 0.08,
        "Пристав / Должностное лицо": 0.04, "Суд-основание": 0.06, "Номер судебного дела": 0.08,
        "Взыскатель": 0.10, "Реквизиты взыскателя": 0.06, "Должник": 0.10, "Реквизиты должника": 0.06,
        "Предмет исполнения": 0.07, "Основной долг": 0.07, "Судебные расходы": 0.03,
        "Итоговая сумма и Баланс": 0.07,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": doc.get("file_name", ""),
        "mode": "autonomous",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": bool(is_bal and date_score >= 70.0)
    }


def evaluate_salary_document_autonomous(doc: Dict[str, Any]) -> Dict[str, Any]:
    auth = doc.get("authority") or {}
    emp = doc.get("employer") or {}
    f = doc.get("finances") or {}

    _, date_score = validate_date_string(doc.get("doc_date"))
    _, ip_score = validate_ip_number_format(doc.get("ip_number"))
    _, rate_score = validate_deduction_rate(f.get("deduction_percentage"))
    _, _, db_inn_score = validate_inn_string(doc.get("debtor_details"))
    is_bal, _, bal_score = validate_math_balance(
        f.get("debt_amount_rub"), f.get("fee_penalty_rub"), None, None, f.get("total_deduction_rub")
    )

    scores = {
        "Дата постановления": date_score,
        "Номер ИП": ip_score,
        "Орган исполнения": 100.0 if len(str(auth.get("name", "")).strip()) > 3 else 0.0,
        "Юрисдикция": 100.0 if auth.get("jurisdiction") else 70.0,
        "Работодатель": 100.0 if len(str(emp.get("name", "")).strip()) > 2 else 0.0,
        "Адрес работодателя": 100.0 if emp.get("address") else 60.0,
        "Взыскатель": 100.0 if len(str(doc.get("claimant_name", "")).strip()) > 2 else 0.0,
        "Должник (Сотрудник)": 100.0 if len(str(doc.get("debtor_name", "")).strip()) > 2 else 0.0,
        "Реквизиты должника": db_inn_score,
        "Предмет взыскания": 100.0 if doc.get("claim_subject") else 50.0,
        "Процент удержания (229-ФЗ)": rate_score,
        "Основной долг": 100.0 if f.get("debt_amount_rub") is not None else 80.0,
        "Исполнительский сбор": 100.0 if f.get("fee_penalty_rub") is not None else 90.0,
        "Итого к удержанию и Баланс": bal_score,
        "Первичный судебный акт": 100.0 if doc.get("base_doc") else 60.0,
    }

    weights = {
        "Дата постановления": 0.05, "Номер ИП": 0.10, "Орган исполнения": 0.08, "Юрисдикция": 0.04,
        "Работодатель": 0.08, "Адрес работодателя": 0.04, "Взыскатель": 0.10, "Должник (Сотрудник)": 0.10,
        "Реквизиты должника": 0.05, "Предмет взыскания": 0.06, "Процент удержания (229-ФЗ)": 0.10,
        "Основной долг": 0.06, "Исполнительский сбор": 0.04, "Итого к удержанию и Баланс": 0.06,
        "Первичный судебный акт": 0.05,
    }

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores)
    return {
        "file_name": doc.get("file_name", ""),
        "mode": "autonomous",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": bool(is_bal and rate_score >= 80.0)
    }


def _get_nested(d: Dict[str, Any], path: str) -> Any:
    """Извлечение вложенного значения по точечному пути (например 'finances.total_rub')."""
    if not isinstance(d, dict) or not path:
        return None
    parts = path.split(".")
    curr: Any = d
    for p in parts:
        if isinstance(curr, dict):
            curr = curr.get(p)
        else:
            return None
    return curr


def evaluate_generic_benchmark(
    pred: Dict[str, Any],
    gt: Dict[str, Any],
    doc_type: Optional[str] = None
) -> Dict[str, Any]:
    """
    Универсальная оценка в режиме Benchmark (vs Ground Truth) на основе benchmark.json плагина.
    (Устраняет дефект C-06: исключает ложный fallback на executive_documents).
    """
    scores: Dict[str, float] = {}
    weights: Dict[str, float] = {}

    plugin = None
    if doc_type:
        from ..type_registry import get_registry
        try:
            plugin = get_registry().get(doc_type)
        except Exception as e:
            logger.debug(f"Не удалось загрузить плагин '{doc_type}' для оценки: {e}")

    bench_fields = plugin.benchmark_config.get("fields", []) if plugin else []

    if bench_fields:
        for f in bench_fields:
            path = f.get("path", "")
            if not path:
                continue
            weight = float(f.get("weight", 1.0))
            m_type = f.get("type", "fuzzy")
            label = f.get("label", path)

            val_ext = _get_nested(pred, path)
            val_gt = _get_nested(gt, path)

            if val_ext is None and val_gt is None:
                sim = 100.0
            elif val_ext is None or val_gt is None:
                sim = 0.0
            elif m_type == "exact":
                sim = exact_match_score(val_ext, val_gt)
            elif m_type == "numeric":
                sim = numeric_proximity_score(val_ext, val_gt)
            else:
                sim = text_similarity_score(val_ext, val_gt)

            scores[label] = round(sim, 2)
            weights[label] = weight
    else:
        # Fallback при отсутствии конфига: динамическое попарное сравнение ключей эталона
        for k, gt_val in gt.items():
            if k in ("file_name", "doc_type"):
                continue
            pred_val = pred.get(k)
            if isinstance(gt_val, (int, float)):
                sim = numeric_proximity_score(pred_val, gt_val)
            elif isinstance(gt_val, dict):
                sim = 100.0 if exact_match_score(pred_val, gt_val) == 100.0 else text_similarity_score(str(pred_val), str(gt_val))
            else:
                sim = text_similarity_score(pred_val, gt_val)
            scores[k] = round(sim, 2)
            weights[k] = 1.0

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores) if total_weight > 0 else 100.0

    return {
        "file_name": gt.get("file_name", pred.get("file_name", "")),
        "mode": "benchmark",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": bool(weighted_total >= 70.0)
    }


def evaluate_generic_autonomous(
    doc: Dict[str, Any],
    doc_type: Optional[str] = None
) -> Dict[str, Any]:
    """
    Универсальная автономная оценка качества на основе autonomous.json плагина.
    (Устраняет дефект C-06: проверяет поля соответствующего типа документа).
    """
    scores: Dict[str, float] = {}
    weights: Dict[str, float] = {}

    plugin = None
    if doc_type:
        from ..type_registry import get_registry
        try:
            plugin = get_registry().get(doc_type)
        except Exception as e:
            logger.debug(f"Не удалось загрузить плагин '{doc_type}' для оценки: {e}")

    rules = plugin.autonomous_config.get("fields", []) if plugin else []

    if rules:
        for r in rules:
            f_path = r.get("field", "")
            if not f_path:
                continue
            rule = r.get("rule", "required")
            val = _get_nested(doc, f_path)

            score = 0.0
            if rule in ("required", "not_empty"):
                score = 100.0 if val not in (None, "", [], {}) else 0.0
            elif rule == "date_format":
                _, score = validate_date_string(val)
            elif rule == "positive_number":
                try:
                    num = float(str(val).replace(" ", "").replace(",", "."))
                    score = 100.0 if num > 0 else 0.0
                except (ValueError, TypeError):
                    score = 0.0
            elif rule == "valid_inn":
                _, _, score = validate_inn_string(val)
            else:
                score = 100.0 if val not in (None, "", [], {}) else 0.0

            scores[f_path] = score
            weights[f_path] = 1.0
    else:
        # Completeness fallback: проверка заполненности ключей
        for k, v in doc.items():
            if k in ("file_name", "doc_type"):
                continue
            scores[k] = 100.0 if v not in (None, "", [], {}) else 0.0
            weights[k] = 1.0

    total_weight = sum(weights.values())
    weighted_total = sum((scores[k] * (weights[k] / total_weight)) for k in scores) if total_weight > 0 else 100.0

    return {
        "file_name": doc.get("file_name", ""),
        "mode": "autonomous",
        "overall_score": round(weighted_total, 2),
        "field_scores": scores,
        "validation_passed": bool(weighted_total >= 70.0)
    }


BENCHMARK_EVALUATORS = {
    "executive": evaluate_executive_document_benchmark,
    "executive_documents": evaluate_executive_document_benchmark,
    "enforcement": evaluate_enforcement_document_benchmark,
    "enforcement_orders": evaluate_enforcement_document_benchmark,
    "salary": evaluate_salary_document_benchmark,
    "salary_deductions": evaluate_salary_document_benchmark,
    "commercial_contracts": lambda p, g: evaluate_generic_benchmark(p, g, "commercial_contracts"),
    "invoices_upd": lambda p, g: evaluate_generic_benchmark(p, g, "invoices_upd"),
    "acceptance_certificates": lambda p, g: evaluate_generic_benchmark(p, g, "acceptance_certificates"),
    "powers_of_attorney": lambda p, g: evaluate_generic_benchmark(p, g, "powers_of_attorney"),
    "legal_claims": lambda p, g: evaluate_generic_benchmark(p, g, "legal_claims"),
    "hr_orders": lambda p, g: evaluate_generic_benchmark(p, g, "hr_orders"),
}

AUTONOMOUS_EVALUATORS = {
    "executive": evaluate_executive_document_autonomous,
    "executive_documents": evaluate_executive_document_autonomous,
    "enforcement": evaluate_enforcement_document_autonomous,
    "enforcement_orders": evaluate_enforcement_document_autonomous,
    "salary": evaluate_salary_document_autonomous,
    "salary_deductions": evaluate_salary_document_autonomous,
    "commercial_contracts": lambda d: evaluate_generic_autonomous(d, "commercial_contracts"),
    "invoices_upd": lambda d: evaluate_generic_autonomous(d, "invoices_upd"),
    "acceptance_certificates": lambda d: evaluate_generic_autonomous(d, "acceptance_certificates"),
    "powers_of_attorney": lambda d: evaluate_generic_autonomous(d, "powers_of_attorney"),
    "legal_claims": lambda d: evaluate_generic_autonomous(d, "legal_claims"),
    "hr_orders": lambda d: evaluate_generic_autonomous(d, "hr_orders"),
}


def get_benchmark_evaluator(doc_type: Optional[str] = None):
    """Возвращает оценщик бенчмарка для типа документа с безопасным fallback."""
    if doc_type and doc_type in BENCHMARK_EVALUATORS:
        return BENCHMARK_EVALUATORS[doc_type]
    return lambda p, g: evaluate_generic_benchmark(p, g, doc_type=doc_type)


def get_autonomous_evaluator(doc_type: Optional[str] = None):
    """Возвращает автономный оценщик для типа документа с безопасным fallback."""
    if doc_type and doc_type in AUTONOMOUS_EVALUATORS:
        return AUTONOMOUS_EVALUATORS[doc_type]
    return lambda d: evaluate_generic_autonomous(d, doc_type=doc_type)


def get_status_from_score(score: float) -> str:
    """Определяет категорию качества по шкале 0..100%."""
    if score >= 95.0:
        return "excellent"
    elif score >= 85.0:
        return "high"
    elif score >= 70.0:
        return "satisfactory"
    return "needs_attention"


# ==============================================================================
# 5. ГЛАВНЫЙ МЕНЕДЖЕР ОЦЕНКИ КАТЕГОРИИ
# ==============================================================================

def evaluate_dataset(
    predicted_docs: Any,
    ground_truth_docs: Optional[Any] = None,
    doc_type: Optional[Any] = None,
    registry: Optional[Any] = None
) -> Any:
    """
    Универсальная оценка датасета:
    - Поддерживает как списки документов (List[Dict]), так и пути к каталогам (str, Path) для MCP и CLI (C-08).
    - Если передан ground_truth_docs: сопоставляет по имени файла и рассчитывает Benchmark Quality Score (%).
    - Если ground_truth_docs не передан или файл в нем отсутствует: выполняет автономную оценку Autonomous Quality & Guardrails (%).
    - Для всех 9 типов документов использует целевые правила плагинов, исключая ложный fallback на executive_documents (C-06).
    """
    from pathlib import Path
    import glob

    # Разрешение параметров при вызове из MCP: evaluate_dataset(results_dir, gt_dir, registry)
    if registry is None and doc_type is not None and not isinstance(doc_type, str):
        registry = doc_type
        doc_type = None

    # Обработка передачи путей каталогов (MCP / Batch mode)
    if isinstance(predicted_docs, (str, Path)) and os.path.isdir(str(predicted_docs)):
        results_dir = Path(predicted_docs)
        gt_dir = Path(ground_truth_docs) if ground_truth_docs and os.path.isdir(str(ground_truth_docs)) else None

        results_map: Dict[str, List[Dict[str, Any]]] = {}
        for jf in glob.glob(str(results_dir / "*.json")):
            base = os.path.basename(jf)
            if base.startswith("run_metrics") or base.startswith("benchmark_") or base.endswith("_quality_metrics.json"):
                continue
            try:
                with open(jf, "r", encoding="utf-8") as f:
                    d = json.load(f)
                    items = d if isinstance(d, list) else [d]
                    for it in items:
                        if isinstance(it, dict):
                            dt = it.get("doc_type", "unknown")
                            results_map.setdefault(dt, []).append(it)
            except Exception as e:
                logger.debug(f"Пропущен некорректный JSON файл '{jf}': {e}")

        if registry is None:
            from ..type_registry import get_registry
            registry = get_registry()

        category_metrics: Dict[str, Dict[str, Any]] = {}
        all_doc_types = set(results_map.keys())
        if registry and hasattr(registry, "enabled"):
            all_doc_types.update(registry.enabled().keys())

        for cat_id in all_doc_types:
            if doc_type and cat_id != doc_type:
                continue
            cat_preds = results_map.get(cat_id, [])
            cat_gt = None
            cat_plugin = registry.get(cat_id) if registry is not None else None
            if gt_dir and registry and hasattr(registry, "enabled") and cat_id in registry.enabled() and cat_plugin is not None:
                gt_file = cat_plugin.gt_file
                gt_fpath = gt_dir / gt_file
                if gt_fpath.is_file():
                    try:
                        with open(gt_fpath, "r", encoding="utf-8") as gf:
                            cat_gt = json.load(gf)
                    except Exception as e:
                        logger.debug(f"Не удалось прочитать эталон '{gt_fpath}': {e}")
                        cat_gt = None
            elif gt_dir:
                candidate = gt_dir / f"{cat_id}.json"
                if candidate.is_file():
                    try:
                        with open(candidate, "r", encoding="utf-8") as gf:
                            cat_gt = json.load(gf)
                    except Exception as e:
                        logger.debug(f"Не удалось прочитать эталон '{candidate}': {e}")
                        cat_gt = None

            if cat_preds or cat_gt:
                category_metrics[cat_id] = evaluate_dataset(cat_preds, cat_gt, doc_type=cat_id)

        if doc_type and doc_type in category_metrics:
            return category_metrics[doc_type]
        return category_metrics

    # Режим списка документов
    doc_list: List[Dict[str, Any]] = predicted_docs if isinstance(predicted_docs, list) else []

    doc_type_str: str
    if not doc_type and doc_list:
        first_doc = doc_list[0]
        first_data: Any = first_doc.get("data")
        p_check: Dict[str, Any] = first_data if isinstance(first_data, dict) else first_doc
        doc_type_str = str(p_check.get("doc_type") or first_doc.get("doc_type") or "generic")
    elif doc_type:
        doc_type_str = str(doc_type)
    else:
        doc_type_str = "generic"

    bench_func = get_benchmark_evaluator(doc_type_str)
    auto_func = get_autonomous_evaluator(doc_type_str)

    gt_map: Dict[str, Dict[str, Any]] = {}
    if ground_truth_docs and isinstance(ground_truth_docs, list):
        for d in ground_truth_docs:
            if isinstance(d, dict) and "file_name" in d:
                gt_map[d["file_name"]] = d

    evaluated_docs = []
    status_counts = {"excellent": 0, "high": 0, "satisfactory": 0, "needs_attention": 0}
    total_scores = []

    for raw_p in doc_list:
        if not isinstance(raw_p, dict):
            continue
        p_val: Any = raw_p.get("data")
        p: Dict[str, Any] = p_val if isinstance(p_val, dict) else raw_p
        f_name = p.get("file_name") or raw_p.get("file_name", "")

        if f_name in gt_map:
            doc_res = bench_func(p, gt_map[f_name])
        else:
            doc_res = auto_func(p)

        score = doc_res["overall_score"]
        status = get_status_from_score(score)
        doc_res["status"] = status
        doc_res["file_name"] = f_name

        status_counts[status] += 1
        total_scores.append(score)
        evaluated_docs.append(doc_res)

    avg_score = round(sum(total_scores) / len(total_scores), 2) if total_scores else 0.0

    return {
        "doc_type": doc_type_str,
        "total_documents": len(evaluated_docs),
        "mode": "benchmark" if gt_map else "autonomous",
        "average_quality_score_percent": avg_score,
        "status_counts": status_counts,
        "documents": evaluated_docs
    }


def generate_run_summary(all_category_metrics: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """Формирует общую сводку запуска (Run Summary) по всем категориям документов."""
    # Защита: если передан отчет по одной категории напрямую (C-08)
    if isinstance(all_category_metrics, dict) and "documents" in all_category_metrics and "average_quality_score_percent" in all_category_metrics:
        dt = str(all_category_metrics.get("doc_type", "default"))
        wrapped: Dict[str, Dict[str, Any]] = {dt: all_category_metrics}
        all_category_metrics = wrapped

    total_docs = 0
    all_scores = []
    status_dist = {"excellent": 0, "high": 0, "satisfactory": 0, "needs_attention": 0}
    categories_summary = {}

    for cat_name, metrics in all_category_metrics.items():
        if not isinstance(metrics, dict):
            continue
        doc_count = metrics.get("total_documents", 0)
        total_docs += doc_count
        c_status = metrics.get("status_counts", {})
        for k in status_dist:
            status_dist[k] += c_status.get(k, 0)

        for d in metrics.get("documents", []):
            all_scores.append(d.get("overall_score", 0.0))

        categories_summary[cat_name] = {
            "total_documents": doc_count,
            "average_quality_score_percent": metrics.get("average_quality_score_percent", 0.0),
            "mode": metrics.get("mode", "autonomous"),
            "status_counts": c_status
        }

    overall_score = round(sum(all_scores) / len(all_scores), 2) if all_scores else 0.0

    return {
        "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_documents_processed": total_docs,
        "overall_quality_score_percent": overall_score,
        "status_distribution": status_dist,
        "categories": categories_summary
    }


# ==============================================================================
# 6. ЭКСПОРТ ОТЧЕТОВ: JSON, MARKDOWN, EXCEL
# ==============================================================================

def export_metrics_json(data: Dict[str, Any], filepath: str):
    """Экспортирует метрический отчет в JSON файл (атомарно, M-19)."""
    dir_path = os.path.dirname(os.path.abspath(filepath))
    os.makedirs(dir_path, exist_ok=True)
    write_atomic(filepath, json.dumps(data, ensure_ascii=False, indent=2))


def append_to_metrics_history(run_summary: Dict[str, Any], history_filepath: str):
    """Добавляет сводку запуска в накопительную историю metrics_history.json (атомарно, M-19)."""
    dir_path = os.path.dirname(os.path.abspath(history_filepath))
    os.makedirs(dir_path, exist_ok=True)

    history = []
    if os.path.exists(history_filepath):
        try:
            with open(history_filepath, "r", encoding="utf-8") as f:
                history = json.load(f)
                if not isinstance(history, list):
                    history = []
        except Exception as e:
            logger.warning(f"Не удалось прочитать историю метрик '{history_filepath}' (перезаписывается): {e}")
            history = []

    history.append(run_summary)
    write_atomic(history_filepath, json.dumps(history, ensure_ascii=False, indent=2))


def export_run_summary_markdown(summary: Dict[str, Any], filepath: str):
    """Экспортирует читаемый Markdown-отчет о качестве запуска."""
    dir_path = os.path.dirname(os.path.abspath(filepath))
    os.makedirs(dir_path, exist_ok=True)

    score = summary.get("overall_quality_score_percent", 0.0)
    total = summary.get("total_documents_processed", 0)
    st = summary.get("status_distribution", {})

    md = [
        "# 📊 Сводный отчет о качестве распознавания документов (Run Metrics)",
        f"> **Дата запуска:** {summary.get('timestamp')}  ",
        f"> **Обработано документов:** {total}  ",
        f"> **Итоговый Quality Score:** **{score}%**  \n",
        "## 1. Распределение качества документов",
        f"- 🟢 **Отличное качество (>= 95%):** {st.get('excellent', 0)}",
        f"- 🟡 **Высокое качество (85-94%):** {st.get('high', 0)}",
        f"- 🟠 **Удовлетворительное (70-84%):** {st.get('satisfactory', 0)}",
        f"- 🔴 **Требует внимания (< 70%):** {st.get('needs_attention', 0)}\n",
        "## 2. Результаты по категориям",
        "| Категория | Документов | Режим оценки | Quality Score (%) |",
        "| :--- | :---: | :---: | :---: |"
    ]

    for cat_k, cat_v in summary.get("categories", {}).items():
        md.append(f"| `{cat_k}` | {cat_v.get('total_documents')} | {cat_v.get('mode')} | **{cat_v.get('average_quality_score_percent')}%** |")

    md.append("\n---\n*Отчет сгенерирован автоматически AI-Системой ScanReader.*")

    write_atomic(filepath, "\n".join(md))


def export_run_summary_excel(summary: Dict[str, Any], filepath: str):
    """Экспортирует стилизованную Excel-сводку о запуске."""
    if not OPENPYXL_AVAILABLE:
        return

    dir_path = os.path.dirname(os.path.abspath(filepath))
    os.makedirs(dir_path, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка запуска"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    headers = ["Категория документов", "Количество", "Режим оценки", "Quality Score (%)", "🟢 Отлично", "🟡 Высокое", "🟠 Удовл.", "🔴 Внимание"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for cat_k, cat_v in summary.get("categories", {}).items():
        st = cat_v.get("status_counts", {})
        ws.append([
            cat_k,
            cat_v.get("total_documents", 0),
            cat_v.get("mode", ""),
            cat_v.get("average_quality_score_percent", 0.0),
            st.get("excellent", 0),
            st.get("high", 0),
            st.get("satisfactory", 0),
            st.get("needs_attention", 0)
        ])

    wb.save(filepath)
